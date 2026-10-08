from pathlib import Path
from openpyxl import load_workbook
import csv,json
b=Path(__file__).resolve().parent
w=load_workbook(b/'rosstat_regional_monthly_wage_snapshot_2026.xlsx',read_only=True,data_only=True)
s=w['с 2019'];r=list(s.values);year=None;cols=[]
for c in range(1,len(r[1])):
 if r[1][c]:year=int(str(r[1][c])[:4])
 if year in [2023,2024]:cols.append((c,year,(c-1)%12+1))
rows=[]
for c,y,m in cols:
 value=r[3][c];assert isinstance(value,(int,float))
 rows.append({'date':f'{y}-{m:02}','value':value,'units':'nominal_rub_per_employee_month','scope':'Российская Федерация','source_url':'https://rosstat.gov.ru/storage/mediabank/tab2-zpl_07-2026.xlsx','source_sheet':'с 2019','source_cell':s.cell(4,c+1).coordinate,'snapshot_release_date':'2026-09-30','published_at':'','available_from':'','historical_availability':'unknown','use_gate':'exploratory_snapshot_only_hypothetical_lag_not_asof_proof'})
assert len(rows)==24 and len(set(x['date'] for x in rows))==24
with (b/'rosstat_national_monthly_wage_2023_2024_snapshot2026.csv').open('w') as f:
 wr=csv.DictWriter(f,fieldnames=list(rows[0]));wr.writeheader();wr.writerows(rows)
print(rows[0],rows[-1])
